"""八幡地区_検証試作_v01.xlsx の動作確認（試験専用コピーで実施）。

- 納品用ファイルは変更しない。試験用の架空入力は output/試験専用/ のコピーにだけ入れる。
- 再計算は LibreOffice（recalc.py）で行う。Excel実機での確認ではない。
- 結果は output/試験専用/試験結果_v01.json と標準出力へ書き出す（試作確認結果_v01.md の表の元）。

使い方: RECALC=/path/to/recalc.py python3 scripts/03_試作検証.py
"""
import datetime as dt
import glob
import json
import os
import shutil
import subprocess
import zipfile

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "output", "八幡地区_検証試作_v01.xlsx")
TDIR = os.path.join(ROOT, "output", "試験専用")
RECALC = os.environ.get("RECALC") or glob.glob("/root/.claude/skills/**/xlsx/scripts/recalc.py", recursive=True)[0]
TARGET = dt.date(2025, 12, 6)
results = []


def recalc(path):
    out = subprocess.run(["python3", RECALC, path, "300"], capture_output=True, text=True)
    info = json.loads(out.stdout)
    assert info.get("status") == "success" and info["total_errors"] == 0, info
    return info


def check(no, name, inputs, expected, actual, method="LibreOffice再計算"):
    ok = expected == actual
    results.append({"no": no, "name": name, "inputs": inputs, "expected": expected, "actual": actual,
                    "result": "合格" if ok else "不合格", "method": method})
    print(("OK " if ok else "NG ") + no, name, "期待", expected, "実際", actual)


def hhmm(v):
    if isinstance(v, dt.datetime):
        v = v.time()
    return v.strftime("%H:%M") if isinstance(v, dt.time) else v


# ---------------------------------------------------------------- 元帳からの独立計算（Excelの式を使わない）
def independent():
    path = glob.glob(os.path.join(ROOT, "input", "01_元資料", "作業実績*"))[0]
    ws = openpyxl.load_workbook(path, read_only=True)["R7年度"]
    iv, jobs, job = [], 0, 0
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1 or r[0] is None or r[17] is None:
            continue
        if dt.date(2025 if r[0] >= 4 else 2026, r[0], r[1]) != TARGET:
            continue
        if r[4] == 1:
            jobs += 1
            job = jobs
        s, e = r[11], r[12]
        sm, em = s.hour * 60 + s.minute, e.hour * 60 + e.minute
        if em < sm:
            em += 1440
        iv.append((sm, em, r[17], job))
    best, at, njobs, m3 = 0, None, 0, 0
    for t in range(0, 1440, 5):
        act = [(x[2], x[3]) for x in iv if x[0] <= t < x[1]]
        n = len({a for a, _ in act})
        if n >= 3:
            m3 += 5
        if n > best:
            best, at, njobs = n, t, len({b for _, b in act})
    return {"rows": len(iv), "jobs": jobs, "max": best, "at": f"{at // 60:02d}:{at % 60:02d}", "jobs_at": njobs, "min3": m3}


# ---------------------------------------------------------------- プログラム検査（構造）
def structure_checks():
    wb = openpyxl.load_workbook(SRC)
    names = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    check("P1", "シート名が作成条件どおり", "—",
          ["はじめに", "実績", "前日確定", "変更・キャンセル", "入渠", "他港作業", "他社配置", "標準時間", "マスター", "集計"],
          names, "プログラム検査")
    with zipfile.ZipFile(SRC) as z:
        macro = any("vbaProject" in n for n in z.namelist())
    check("P2", "マクロ・VBAを含まない", "—", False, macro, "プログラム検査")
    bad = []
    for ws in wb.worksheets:
        if ws.title == "はじめに":
            continue
        for row in ws.iter_rows(min_row=5, max_row=25):
            for c in row:
                v = c.value
                if isinstance(v, str) and ("【入力" in v or "【自動" in v) and "／" not in v:
                    if c.column == 1 and ws.title == "前日確定" and (6 <= c.row <= 10 or c.row == 22):  # 左に見出し、C列が値の欄
                        ok = ws.cell(c.row, 3).fill.fgColor.rgb[-6:] == ("FFF9C4" if "【入力" in v else "EDEDED")
                    else:
                        ok = c.fill.fgColor.rgb[-6:] == ("FFE97F" if "【入力" in v else "BFBFBF")
                    if not ok:
                        bad.append(f"{ws.title}!{c.coordinate}")
    check("P3", "見出しの【入力】は黄色・【自動】は灰色", "—", [], bad, "プログラム検査")
    dv = {ws.title: sorted(str(d.sqref) for d in ws.data_validations.dataValidation) for ws in wb.worksheets}
    need = {"前日確定": ["B14:B21", "C10", f"K26:P45", "S26:S45"],
            "変更・キャンセル": ["A26:A85", "L26:Q85", "R26:R85", "A15:A22", "D15:D22", "C7:C11"],
            "入渠": ["A6:A55"], "他港作業": ["A6:A55"], "他社配置": ["A6:A55", "B6:B55"], "マスター": ["C6:C125"]}
    missing = [f"{s}!{r}" for s, rs in need.items() for r in rs if r not in dv[s]]
    check("P4", "プルダウン（船名・作業・○）が入力欄に設定されている", "—", [], missing, "プログラム検査")
    sel = (wb.active.title, wb["はじめに"]["B17"].hyperlink.location, wb["前日確定"].sheet_view.selection[0].activeCell)
    check("P5", "初期表示：はじめにのリンクから最初の入力欄（前日確定C7）へ案内", "—",
          ("はじめに", "'前日確定'!C7", "C7"), sel, "プログラム検査")


# ---------------------------------------------------------------- 試験コピーA：入力→集計のつながり
def label_yotei(n, time, io, ship):
    return f"{TARGET:%m/%d} 予定のみ#{n:02d} {time} {io} {ship}"


def copy_a():
    path = os.path.join(TDIR, "試験A_入力と集計のつながり.xlsx")
    shutil.copy(SRC, path)
    wb = openpyxl.load_workbook(path)
    z, h, m = wb["前日確定"], wb["変更・キャンセル"], wb["マスター"]
    lab07 = z["B32"].value  # 12/06 #07 LNG MARS
    lab01 = z["B26"].value
    assert "LNG MARS" in lab07
    # 前日確定（架空）
    z["C7"] = "【試験】架空の予定表"
    z["J32"] = dt.time(13, 0)
    for c, v in zip("KLMNO", ["S", "T", "M", "P", "Y"]):
        z[f"{c}32"] = v
    z["J31"] = dt.time(12, 30)
    z["K31"], z["L31"] = "竜山", "さくら"
    z["D35"], z["E35"], z["F35"], z["J35"] = "試験船X", "入", "N4", dt.time(10, 0)
    for r, v in zip(range(14, 18), ["T", "S", "T", "試験タグA"]):
        z[f"B{r}"] = v
    # マスター追加・修正（架空）
    m["A40"], m["C40"], m["D40"], m["E40"], m["I40"] = "試験タグA", "試験会社", "タグ", "門司港", "確認済み"
    m["P15"], m["Q15"] = "試験会社", "他社"
    m["B11"] = "鐵豊丸"          # 記号T＝鐵豊丸と仮定（試験のみ）
    m["C29"] = "洞海マリン"      # さくら＝洞海マリンと仮定（試験のみ）
    m["Z1"] = "=ROWS(船名リスト)"
    # 変更・キャンセル（架空）
    h["B7"] = "変更あり（下の表に入力）"
    h["A8"], h["C8"] = dt.date(2025, 12, 5), "○"
    h["A26"], h["G26"], h["J26"], h["V26"] = lab07, dt.datetime(2025, 12, 6, 8, 0), dt.time(14, 0), "あり"
    h["A27"], h["R27"], h["S27"], h["T27"] = label_yotei(1, "10:00", "入", "試験船X"), "○", "当日・基地発後", "本船都合"
    h["A28"], h["G28"], h["J28"] = lab07, dt.datetime(2025, 12, 6, 11, 0), dt.time(15, 30)
    for c, v in zip("LMNO", ["S", "T", "M", "P"]):
        h[f"{c}28"] = v
    calls = [("T", dt.datetime(2025, 12, 6, 6, 0), "乗組員の出勤手配の不足"),
             ("T", dt.datetime(2025, 12, 6, 6, 0), "乗組員の出勤手配の不足"),
             ("関", dt.datetime(2025, 12, 6, 6, 30), "常駐船が他作業で不足"),
             ("竜山", dt.datetime(2025, 12, 6, 7, 0), "その他")]
    for n, (s, t, why) in enumerate(calls):
        r = 15 + n
        h[f"A{r}"], h[f"C{r}"], h[f"D{r}"], h[f"E{r}"] = s, t, lab01, why
    # 入渠・他港作業・他社配置（架空）
    wb["入渠"]["A6"], wb["入渠"]["C6"], wb["入渠"]["D6"] = "Y", dt.date(2025, 12, 1), dt.date(2025, 12, 10)
    wb["入渠"]["A7"], wb["入渠"]["C7"], wb["入渠"]["D7"] = "M", dt.date(2025, 12, 7), dt.date(2025, 12, 20)
    o = wb["他港作業"]
    o["A6"], o["C6"], o["D6"], o["E6"] = "関", "試験港", dt.datetime(2025, 12, 6, 5, 0), dt.datetime(2025, 12, 6, 9, 0)
    o["A7"], o["C7"], o["D7"], o["E7"] = "T", "試験港", dt.datetime(2025, 12, 7, 5, 0), dt.datetime(2025, 12, 7, 9, 0)
    t = wb["他社配置"]
    t["A6"], t["B6"], t["C6"], t["D6"], t["E6"] = "洞海マリン", "さくら", "若松", dt.date(2025, 4, 1), dt.date(2026, 3, 31)
    wb.save(path)
    recalc(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    s, z, h, m, j = wb["集計"], wb["前日確定"], wb["変更・キャンセル"], wb["マスター"], wb["実績"]
    check("A1", "前日確定の入力が集計へ反映（予定時刻入力作業数・予定タグ延べ・予定のみ）",
          "#07 13:00 S,T,M,P,Y／#06 12:30 竜山,さくら／予定のみ 試験船X 10:00", [3, 7, 1, "S・T・M・P・Y"],
          [s["C41"].value, s["C43"].value, s["C42"].value, z["R32"].value])
    check("A2", "翌日の出勤手配隻数（同じ船の重複は1隻）", "T・S・T・試験タグA", [3, "試験会社"],
          [s["C44"].value, z["C17"].value])
    check("A3", "複数回の変更：前日確定を上書きせず、変更前の時刻・タグを順に引き継ぐ",
          "前日13:00 → 1回目14:00 → (間に別作業の行) → 2回目15:30・タグS,T,M,P",
          ["13:00", 1, "13:00", 2, "14:00", "S・T・M・P・Y", "13:00", 2, "15:30"],
          [hhmm(h["E26"].value), h["D26"].value, hhmm(h["I26"].value), h["D28"].value, hhmm(h["I28"].value),
           h["K28"].value, hhmm(z["J32"].value), z["U32"].value, hhmm(z["V32"].value)])
    check("A4", "実績のない取消しを予定表から追加し、取消しとして集計（実績は変わらない）",
          "前日確定35行に試験船X、変更で取消し○・当日・基地発後", [1, 1, 3, 2, 1, 17, 8],
          [s["C53"].value, s["C54"].value, s["C51"].value, s["C52"].value, s["C55"].value, s["C10"].value, s["C14"].value])
    check("A5", "マスターへ船・会社を追加すると選択肢と会社別集計へ反映",
          "船「試験タグA」・会社「試験会社(他社)」を追加", [35, "試験会社"],
          [m["Z1"].value, s["L19"].value])
    check("A6", "正式船名を入れると同じ船として1隻に数える（記号T＝鐵豊丸と仮定）",
          "マスターT行の正式船名＝鐵豊丸", ["鐵豊丸", 8, 5], [j["W8"].value, s["C14"].value, s["D30"].value])
    check("A7", "会社を入れると会社区分・会社別へ反映（さくら＝洞海マリンと仮定）",
          "マスターさくら行の会社＝洞海マリン", [13, 2, 1, 1, 1],
          [s["C30"].value, s["C31"].value, s["C32"].value, s["C33"].value, s["N17"].value])
    check("A8", "休船からの呼出：同じ船の同じ呼出は1回、乗組員の手配不足を区別、対象外の会社を表示",
          "T 6:00×2行（乗組員）、関 6:30、竜山 7:00", [3, 1, 1, "同一呼出（1回と数える）"],
          [s["C56"].value, s["C57"].value, s["C58"].value, h["H16"].value])
    check("A9", "入渠・他港作業・他社配置が対象日に反映",
          "入渠Y 12/1～12/10、入渠M 12/7～、他港 関 12/6、他港 T 12/7、他社配置 さくら", [1, 1, 1, "西日本海運以外の船です（確認してください）", ""],
          [s["C61"].value, s["C62"].value, s["C63"].value, wb["他港作業"]["J7"].value, wb["他社配置"]["J6"].value or ""])
    check("A10", "日ごとの変更確認と全船休航（別の日の休航は対象日に混ざらない）",
          "12/6 変更あり、12/5 全船休航○", ["変更あり（下の表に入力）", None],
          [s["C49"].value, s["C50"].value])


# ---------------------------------------------------------------- 試験コピーB：取込データを試験用に書き換えた確認
def copy_b():
    path = os.path.join(TDIR, "試験B_取込データの書き換え確認.xlsx")
    shutil.copy(SRC, path)
    wb = openpyxl.load_workbook(path)
    j = wb["実績"]
    j["M13"], j["M14"] = "支給炭", "支給炭"        # TEXAS HARMONY を安瀬候補に（試験のみ）
    j["O10"], j["P10"] = dt.time(12, 30), dt.time(13, 0)  # FIRST AI の S を LNG MARS の S と重ねる（試験のみ）
    j["U19"] = "○"                                  # LNG MARS の Y を取消しに（試験のみ）
    wb.save(path)
    recalc(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    s, j = wb["集計"], wb["実績"]
    check("B1", "安瀬候補は応援（他社・会社未確認）と別に数える", "TEXAS HARMONY 2行の品名を支給炭に", [2, 0],
          [s["C35"].value, s["C36"].value])
    check("B2", "同じ船を同じ時間の別作業に入れると重複配船の警告、同時隻数は二重に数えない",
          "FIRST AI(S)を12:30～13:00に変更（LNG MARSのSと重なる）", [2, "同じ船の時間が重なっています", 8],
          [s["C24"].value, j["AH10"].value, s["C14"].value])
    check("B3", "取消しの行は同時稼働に含めない", "LNG MARS(Y)に取消し○", [1, 16, 8],
          [s["C11"].value, s["C13"].value, s["C14"].value])


def main():
    os.makedirs(TDIR, exist_ok=True)
    structure_checks()
    base = os.path.join(TDIR, "試験0_納品ファイルの再計算.xlsx")
    shutil.copy(SRC, base)
    info = recalc(base)
    wb = openpyxl.load_workbook(base, data_only=True)
    s = wb["集計"]
    ind = independent()
    check("R1", "納品ファイルを再計算してエラーがない", "入力なし", 0, info["total_errors"])
    check("R2", "実績の取込（行数・作業数）と同時稼働が、元帳からの独立計算と一致",
          "元帳R7年度 2025/12/6", [ind["rows"], ind["jobs"], ind["max"], ind["at"], ind["jobs_at"], ind["min3"]],
          [s["C10"].value, s["C9"].value, s["C14"].value, hhmm(s["C15"].value), s["C16"].value, s["C17"].value])
    copy_a()
    copy_b()
    with open(os.path.join(TDIR, "試験結果_v01.json"), "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=1, default=str)
    print("合格", sum(r["result"] == "合格" for r in results), "/", len(results))


if __name__ == "__main__":
    main()
